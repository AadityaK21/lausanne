"""
Export Markdown-ish notes to a Microsoft Word (.docx) document.

Handles simple Markdown: '#'/'##'/'###' headings, '-'/'*' bullets, and
inline **bold** runs. Uses python-docx; no external service.
"""

from __future__ import annotations

import re
from pathlib import Path


def _add_runs_with_bold(paragraph, text: str) -> None:
    """Split text on **bold** markers and add styled runs."""
    parts = re.split(r"(\*\*.+?\*\*)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        else:
            paragraph.add_run(part)


def markdown_to_docx(markdown_text: str, out_path: str) -> str:
    from docx import Document

    doc = Document()

    for raw_line in markdown_text.splitlines():
        line = raw_line.rstrip()
        if not line.strip():
            continue

        if line.startswith("### "):
            doc.add_heading(line[4:].strip(), level=3)
        elif line.startswith("## "):
            doc.add_heading(line[3:].strip(), level=2)
        elif line.startswith("# "):
            doc.add_heading(line[2:].strip(), level=1)
        elif line.lstrip().startswith(("- ", "* ")):
            content = line.lstrip()[2:]
            p = doc.add_paragraph(style="List Bullet")
            _add_runs_with_bold(p, content)
        else:
            p = doc.add_paragraph()
            _add_runs_with_bold(p, line)

    Path(out_path).parent.mkdir(parents=True, exist_ok=True)
    doc.save(out_path)
    return out_path
